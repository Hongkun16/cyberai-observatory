#!/usr/bin/env python3
"""Export the reviewed scores alongside the immutable original Excel snapshot."""
from __future__ import annotations

import hashlib
import io
import json
import os
import zipfile

# Keep XML output identical on machines with and without optional lxml installed.
os.environ["OPENPYXL_LXML"] = "False"
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from import_workbook import ROOT, SCORE_REVIEW_DATE, SCORE_STATUS_LABELS, fingerprint, write_json


def export():
    data_path = ROOT / "web/public/data/audit.json"
    data = json.loads(data_path.read_text(encoding="utf-8"))
    workbook = openpyxl.load_workbook(ROOT / data["meta"]["workbook"])
    identities = {m["name"]: m["id"] for m in data["models"]}
    registry = json.loads((ROOT / "data/record-identities.json").read_text())
    records = {r["id"]: r for r in data["evaluations"]}
    extra_headers = ["成绩状态", "成绩复核备注", "成绩证据PDF页码", "记录ID", "原始成绩栏"]
    for sheet_name in ["能力评测", "防护与拒答", "待核对与排除"]:
        sheet = workbook[sheet_name]
        for column, label in enumerate(extra_headers, 15):
            cell = sheet.cell(4, column, label)
            cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="243E63")
            sheet.column_dimensions[cell.column_letter].width = 30 if column != 16 else 65
        for row in range(5, sheet.max_row + 1):
            model = identities[sheet.cell(row, 2).value]
            key = fingerprint(model, sheet.cell(row, 3).value, sheet.cell(row, 5).value, sheet.cell(row, 11).value)
            record = records[registry[key]]
            review = record["scoreAudit"]
            sheet.cell(row, 9, record["score"] or SCORE_STATUS_LABELS[record["scoreStatus"]])
            sheet.cell(row, 8, record["configuration"])
            values = [SCORE_STATUS_LABELS[record["scoreStatus"]], review["notes"] if review else record["status"],
                      "；".join(f"{e['pdfPath']} p.{e['pages']}" for e in review["evidence"]) if review else "",
                      record["id"], record["originalScore"]]
            for column, value in enumerate(values, 15):
                cell = sheet.cell(row, column, value)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.font = Font(name="Microsoft YaHei", size=10, color="344A68")
        sheet.auto_filter.ref = f"A4:S{sheet.max_row}"
        sheet.cell(2, 1, f"原核验快照2026-10-17；PDF成绩补录复核{SCORE_REVIEW_DATE}。数值、图示、定性和不适用分别标注，详见成绩补录审计表。")
    audit = workbook.create_sheet("成绩补录审计")
    audit.append([f"PDF成绩补录审计 | {SCORE_REVIEW_DATE}"])
    audit.append([f"复核{data['meta']['reviewedScoreRecords']}条已确认记录；原成绩{data['meta']['originalScoreCount']}条，现已核实数值{data['meta']['scoreCount']}条，新增数值{data['meta']['addedScoreCount']}条。"])
    audit.append([])
    audit.append(["记录ID", "模型", "评测", "成绩状态", "补录成绩", "原成绩", "复核方式", "证据文件与页码", "原文/图表读数证据", "逐指标数据", "备注"])
    names = {m["id"]: m["name"] for m in data["models"]}
    for r in data["evaluations"]:
        review = r["scoreAudit"]
        audit.append([r["id"], names[r["modelId"]], r["benchmark"], SCORE_STATUS_LABELS[r["scoreStatus"]], r["score"],
                      r["originalScore"], review["method"] if review else "不适用",
                      "；".join(f"{e['pdfPath']} p.{e['pages']}" for e in review["evidence"]) if review else "",
                      "\n".join(e["quote"] for e in review["evidence"]) if review else r["status"],
                      json.dumps(r["metrics"], ensure_ascii=False), review["notes"] if review else r["status"]])
    audit.freeze_panes = "D5"
    audit.auto_filter.ref = f"A4:K{audit.max_row}"
    for row in audit.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.font = Font(name="Microsoft YaHei", size=10, color="344A68")
    for cell in audit[4]:
        cell.fill = PatternFill("solid", fgColor="243E63")
        cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
    for column in "ABCDEFGHIJK":
        audit.column_dimensions[column].width = 45 if column in "EFHIJK" else 24
    instructions = workbook["使用说明"]
    instructions.append(["成绩补录", f"{SCORE_REVIEW_DATE}按PDF正文、表格和图表复核全部196条已确认记录；不从未标值曲线估读精确分数，不将未确认结果归属到模型。"])
    instructions.append(["原始文件", "原2026-10-17工作簿完整保留在仓库snapshots/2026-10-17.xlsx。本文件为附加成绩证据后的导出版本。"])
    # openpyxl/ZIP timestamps would otherwise make reproducible imports look dirty.
    buffer = io.BytesIO()
    workbook.save(buffer)
    output = io.BytesIO()
    with zipfile.ZipFile(buffer) as source, zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as dest:
        for name in sorted(source.namelist()):
            payload = source.read(name)
            if name == "docProps/core.xml":
                import re
                payload = re.sub(rb'(<dcterms:modified[^>]*>).*?(</dcterms:modified>)', rb'\g<1>2026-10-18T00:00:00Z\g<2>', payload)
            info = zipfile.ZipInfo(name, (2026, 10, 18, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            dest.writestr(info, payload)
    destination = ROOT / "web/public" / data["meta"]["downloadWorkbook"]
    destination.write_bytes(output.getvalue())
    data["meta"]["reviewedWorkbookSha256"] = hashlib.sha256(output.getvalue()).hexdigest()
    write_json(ROOT / "data/meta.json", data["meta"])
    write_json(data_path, data)
    print(f"Reviewed workbook exported: {destination.name}")


if __name__ == "__main__":
    export()
