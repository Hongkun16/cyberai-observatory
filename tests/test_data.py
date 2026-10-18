import copy
import hashlib
import importlib.util
import json
import re
import unittest
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location("import_workbook", ROOT / "scripts/import_workbook.py")
importer = importlib.util.module_from_spec(spec)
spec.loader.exec_module(importer)


class AuditIntegrityTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.data = json.loads((ROOT / "web/public/data/audit.json").read_text())
        cls.sheets = importer.read_workbook(ROOT / "snapshots/2026-10-17.xlsx")
        cls.legacy_records = [r for r in cls.data["evaluations"] if r.get("origin") != "vendor-supplement"]
        cls.vendor_records = [r for r in cls.data["evaluations"] if r.get("origin") == "vendor-supplement"]

    def test_workbook_rows_are_not_lost(self):
        actual = Counter((r["modelId"], r["benchmark"], r["status"], r["originalScore"], r["originalConfiguration"], r["pages"]) for r in self.legacy_records)
        models = {m["name"]: m["id"] for m in self.data["models"]}
        expected = Counter((models[r["模型/快照"]], r["Benchmark/评测"], r["证据状态"], r["明确文字数值（非统一指标）"], r["实际执行模型/配置"], r["PDF页码（文件页）"]) for name in ["能力评测", "防护与拒答", "待核对与排除"] for r in self.sheets[name])
        self.assertEqual(actual, expected)
        self.assertEqual(len(self.legacy_records), 233)
        self.assertEqual(len(self.data["evaluations"]), 268)
        self.assertEqual(len(self.data["models"]), 73)

    def test_unconfirmed_records_remain_separate(self):
        self.assertEqual(Counter(r["status"] for r in self.legacy_records), {"已确认": 196, "模型覆盖待核对": 15, "未披露具名项目": 8, "不得归属给Fable生产配置": 8, "明确未运行": 4, "停止报告": 1, "仅方法引用": 1})

    def test_mixed_and_qualified_scores_are_not_converted(self):
        for score in ["plain Cap% 91%；ACE 301/410", ">70%", "100% milestones refused", "10/22 (45.5%)", "2.6×", "7/11至少一次；平均成功率26%", "pass@4 99.2%", "106次control-flow hijack", ""]:
            with self.subTest(score=score):
                self.assertIsNone(importer.parse_percentage(score))
        self.assertEqual(importer.parse_percentage("47.2% pass@1"), 47.2)
        self.assertEqual(importer.parse_percentage("100%"), 100)

    def test_month_precision_does_not_invent_a_day(self):
        monthly = [m for m in self.data["models"] if m["date_precision"] == "month"]
        self.assertEqual(len(monthly), 7)
        self.assertTrue(all(len(m["date"]) == 7 for m in monthly))

    def test_reference_validation_rejects_broken_source_mapping(self):
        importer.validate(self.data)
        broken = copy.deepcopy(self.data)
        broken["evaluations"][0]["pdfPath"] = "pdf/missing.pdf"
        with self.assertRaisesRegex(ValueError, "Unresolved reference"):
            importer.validate(broken)

    def test_rejects_duplicate_ids(self):
        broken = copy.deepcopy(self.data)
        broken["evaluations"][1]["id"] = broken["evaluations"][0]["id"]
        with self.assertRaisesRegex(ValueError, "Duplicate evaluation IDs"):
            importer.validate(broken)

    def test_source_links_and_archive_inventory_preserved(self):
        self.assertEqual(len(self.data["pdfs"]), 69)
        self.assertEqual(sum(p["pages"] for p in self.data["pdfs"]), 4774)
        self.assertEqual(len(self.data["sources"]), 104)
        self.assertTrue(any(s["status"] == "403" for s in self.data["sources"]))
        self.assertTrue(all(len(p["sha256"]) == 64 for p in self.data["pdfs"]))

    def test_no_unreviewed_comparison_groups(self):
        self.assertTrue(all(r["comparisonGroup"] is None for r in self.data["evaluations"]))

    def test_all_confirmed_records_have_archived_score_review(self):
        records = self.data["evaluations"]
        self.assertEqual(sum(r["scoreAudit"] is not None for r in records), 218)
        self.assertEqual(Counter(r["scoreStatus"] for r in records), {"reported": 178, "graphical": 22, "qualitative": 18, "not_applicable": 50})
        self.assertEqual(sum(not r["originalScore"] and r["scoreStatus"] == "reported" for r in self.legacy_records), 113)
        inventory = {p["path"]: p for p in self.data["pdfs"] + self.data["documents"]}
        for r in records:
            if r["scoreAudit"]:
                self.assertTrue(r["scoreAudit"]["evidence"])
                for evidence in r["scoreAudit"]["evidence"]:
                    self.assertEqual(evidence["sha256"], inventory[evidence["pdfPath"]]["sha256"])
                    self.assertTrue(evidence["pages"] or evidence.get("section"))
                    self.assertTrue(evidence["quote"])

    def test_known_pdf_labels_were_transcribed(self):
        records = {r["id"]: r for r in self.data["evaluations"]}
        # Independently checked graph labels and the Cyber Range table, file pages 22/32/33/49.
        for record_id, value in {"CAO-2026-0006": 88.23, "CAO-2026-0019": 93.1, "CAO-2026-0020": 93.33, "CAO-2026-0043": 96.67}.items():
            self.assertEqual(records[record_id]["numericScore"], value)
        self.assertEqual(records["CAO-2026-0033"]["numericScore"], 76.9)
        self.assertIn("13", records["CAO-2026-0033"]["scoreAudit"]["notes"])

    def test_zero_is_a_result_not_a_missing_score(self):
        record = next(r for r in self.data["evaluations"] if r["id"] == "CAO-2026-0081")
        self.assertEqual(record["scoreStatus"], "reported")
        self.assertEqual(record["numericScore"], 0)

    def test_graphical_and_qualitative_results_are_not_estimated(self):
        for r in self.data["evaluations"]:
            if r["scoreStatus"] in ["graphical", "qualitative"]:
                self.assertIsNone(r["numericScore"])
                self.assertEqual(r["metrics"], [])
        r = next(r for r in self.data["evaluations"] if r["id"] == "CAO-2026-0045")
        self.assertEqual(r["scoreStatus"], "graphical")

    def test_source_conflicts_and_later_disclosures_stay_explicit(self):
        records = {r["id"]: r for r in self.data["evaluations"]}
        conflict = records["CAO-2026-0130"]
        self.assertIn("1.0%", conflict["score"])
        self.assertIn("0.9%", conflict["score"])
        self.assertTrue(records["CAO-2026-0198"]["scoreAudit"]["laterDisclosure"])
        self.assertFalse(records["CAO-2026-0112"]["scoreAudit"]["laterDisclosure"])
        self.assertIn("37/40", records["CAO-2026-0112"]["configuration"])
        self.assertIn("3 trials", records["CAO-2026-0129"]["configuration"])

    def test_download_workbook_contains_the_reviewed_scores(self):
        import openpyxl
        path = ROOT / "web/public" / self.data["meta"]["downloadWorkbook"]
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), self.data["meta"]["reviewedWorkbookSha256"])
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        self.assertIn("成绩补录审计", workbook.sheetnames)
        rows = list(workbook["成绩补录审计"].values)[4:]
        self.assertEqual(len(rows), 268)
        self.assertEqual(workbook["新增厂商模型"].max_row, 26)
        self.assertEqual(workbook["新增厂商证据文件"].max_row, 37)
        actual = {r[0]: r for r in rows}
        for record in self.data["evaluations"]:
            self.assertEqual(actual[record["id"]][3], importer.SCORE_STATUS_LABELS[record["scoreStatus"]])
            self.assertEqual(actual[record["id"]][4] or "", record["score"])

    def test_vendor_archives_and_identity_mapping(self):
        supplement = json.loads((ROOT / "data/vendor-supplement.json").read_text())
        registry = json.loads((ROOT / "data/record-identities.json").read_text())
        for item in supplement["documents"] + supplement["provenance"]:
            self.assertEqual(hashlib.sha256((ROOT / item["path"]).read_bytes()).hexdigest(), item["sha256"])
        for r in self.vendor_records:
            self.assertEqual(registry[importer.fingerprint(r["modelId"], r["benchmark"], r["category"], r["sourceUrl"])], r["id"])
        broken = copy.deepcopy(self.data)
        next(r for r in broken["evaluations"] if r.get("origin") == "vendor-supplement" and r["scoreAudit"])["scoreAudit"]["evidence"][0]["sha256"] = "0" * 64
        with self.assertRaisesRegex(ValueError, "Score evidence hash mismatch"):
            importer.validate(broken)

    def test_vendor_snapshots_and_later_disclosures_are_separate(self):
        models = {m["id"]: m for m in self.data["models"]}
        self.assertEqual(Counter(m["provider"] for m in models.values() if m.get("origin") == "vendor-supplement"), {"Tencent": 3, "Zhipu": 8, "DeepSeek": 7, "xAI": 7})
        for mid, date in [("deepseek-v4-flash", "2026-07-31"), ("deepseek-v4-pro", "2026-08-13")]:
            self.assertEqual(models[mid]["date"], date)
            self.assertEqual(models[mid + "-preview"]["date"], "2026-04-24")
        for mid in ["deepseek-v4-flash-preview", "deepseek-v4-pro-preview", "deepseek-v4-1-flash-1017", "grok-4-20-multi-agent", "grok-4-7-fast"]:
            records = [r for r in self.vendor_records if r["modelId"] == mid]
            self.assertTrue(records)
            self.assertTrue(all(r["scoreStatus"] == "not_applicable" and not r["metrics"] for r in records))
        glm5 = [r for r in self.vendor_records if r["modelId"] == "glm-5"]
        self.assertEqual({(r["numericScore"], r["scoreAudit"]["laterDisclosure"]) for r in glm5}, {(43.2, False), (48.3, True)})
        hy3 = next(r for r in self.vendor_records if r["modelId"] == "hy3")
        self.assertTrue(hy3["scoreAudit"]["laterDisclosure"])
        self.assertEqual(hy3["numericScore"], 51.8)

    def test_vendor_units_and_safety_metrics_are_not_conflated(self):
        for mid, values in [("glm-5-2", [29, 39]), ("glm-5-3", [105, 130])]:
            r = next(r for r in self.vendor_records if r["modelId"] == mid and r["benchmark"] == "ExploitGym")
            self.assertIsNone(r["numericScore"])
            self.assertEqual([m["value"] for m in r["metrics"]], values)
            self.assertTrue(all(m["unit"] == "count" for m in r["metrics"]))
            self.assertTrue(r["metrics"][0]["configuration"].endswith("2h预算"))
            self.assertNotIn("6h", r["metrics"][0]["configuration"])
            self.assertTrue(r["metrics"][1]["configuration"].endswith("6h预算"))
            self.assertNotIn("2h", r["metrics"][1]["configuration"])
        safety = next(r for r in self.vendor_records if r["benchmark"] == "HackerBench v0.3")
        self.assertNotEqual(safety["category"], "能力")
        self.assertEqual(safety["metrics"][0]["value"], 3.3)
        grok = {r["benchmark"]: r for r in self.vendor_records if r["modelId"] == "grok-4-20"}
        self.assertEqual({b: r["metrics"][0]["value"] for b, r in grok.items()}, {"Cybench": 53, "WMDP Cyber": 91, "AgentHarm": 30, "AgentDojo": 33})

    def test_benchmark_usage_matches_user_workbook(self):
        import openpyxl
        reference = self.data["benchmarkUsage"]
        path = ROOT / reference["sourceWorkbook"]
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), reference["sha256"])
        rows = list(openpyxl.load_workbook(path, read_only=True, data_only=True).active.values)[1:]
        self.assertEqual([(r[0], r[3]) for r in rows if r[0]], [(r["name"], r["times"]) for r in reference["entries"]])
        self.assertEqual([r["name"] for r in reference["entries"][:6]], ["ExploitBench", "CyberGym", "ExploitGym", "SEC-Bench Pro", "CyScenarioBench", "CVE-Bench"])

    def test_english_translations_cover_displayed_data(self):
        translations = json.loads((ROOT / "web/src/data.en.json").read_text())
        displayed = [m[k] for m in self.data["models"] for k in ["role", "dateNote"]]
        displayed += [b[k] for b in self.data["benchmarks"] for k in ["name", "task", "openness"]]
        displayed += [m[k] for r in self.data["evaluations"] for m in r["metrics"] for k in ["label", "configuration"]]
        for value in set(displayed):
            if re.search(r"[\u3400-\u9fff]", value):
                self.assertIn(value, translations)
                self.assertTrue(translations[value].strip())
                self.assertIsNone(re.search(r"[\u3400-\u9fff]", translations[value]))


if __name__ == "__main__":
    unittest.main()
